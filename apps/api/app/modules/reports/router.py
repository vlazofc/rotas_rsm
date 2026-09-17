import io
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import extract, select
from sqlalchemy.orm import Session, selectinload

from app.core.permissions import FINANCE_VIEW, Role, require_permission, require_roles, scope_by_branch
from app.db.models import (
    Attachment, AuditLog, DeliveryFailureReason, Driver, Expense, FinancialAccount, MaintenanceOrder,
    Part, Revenue, Route, RouteOccurrence, RouteStop, Tire, User, Vehicle, WorkflowTask,
)
from app.db.session import get_db
from app.modules.auth.deps import get_current_user
from app.services import storage

router = APIRouter(prefix="/reports", tags=["reports"])


@router.get("/overview", dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO, Role.AUDITOR))])
def overview_report(start: date | None = None, end: date | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Resumo gerencial multiárea para a central de relatórios."""
    def branch(stmt, model):
        return scope_by_branch(stmt, model.branch_id, user, db)

    route_stmt=branch(select(Route),Route)
    expense_stmt=branch(select(Expense).where(Expense.approval_status=="approved"),Expense)
    revenue_stmt=branch(select(Revenue),Revenue)
    account_stmt=branch(select(FinancialAccount),FinancialAccount)
    if start:
        route_stmt=route_stmt.where(Route.route_date>=start);expense_stmt=expense_stmt.where(Expense.expense_date>=start)
        revenue_stmt=revenue_stmt.where(Revenue.revenue_date>=start);account_stmt=account_stmt.where(FinancialAccount.issue_date>=start)
    if end:
        route_stmt=route_stmt.where(Route.route_date<=end);expense_stmt=expense_stmt.where(Expense.expense_date<=end)
        revenue_stmt=revenue_stmt.where(Revenue.revenue_date<=end);account_stmt=account_stmt.where(FinancialAccount.issue_date<=end)
    routes=list(db.scalars(route_stmt).all());expenses=list(db.scalars(expense_stmt).all());revenues=list(db.scalars(revenue_stmt).all());accounts=list(db.scalars(account_stmt).all())

    def sum_by(rows,key,value):
        result={}
        for row in rows:
            label=str(key(row) or "Não informado");result[label]=round(result.get(label,0)+float(value(row) or 0),2)
        return [{"label":label,"value":value} for label,value in sorted(result.items(),key=lambda item:item[1],reverse=True)]

    vehicle_ids={row.vehicle_id for row in expenses if row.vehicle_id};vehicles={v.id:v.plate for v in db.scalars(select(Vehicle).where(Vehicle.id.in_(vehicle_ids))).all()} if vehicle_ids else {}
    route_status={};
    for row in routes:route_status[row.status]=route_status.get(row.status,0)+1
    revenue_total=round(sum(float(row.amount or 0) for row in revenues),2);expense_total=round(sum(float(row.amount or 0) for row in expenses),2)
    pending_payable=sum(float(row.amount or 0) for row in accounts if row.kind=="payable" and row.status not in {"paga","cancelada"})
    pending_receivable=sum(float(row.amount or 0) for row in accounts if row.kind=="receivable" and row.status not in {"recebida","cancelada"})

    occurrence_stmt=branch(select(RouteOccurrence),RouteOccurrence);task_stmt=branch(select(WorkflowTask),WorkflowTask);maintenance_stmt=branch(select(MaintenanceOrder),MaintenanceOrder);tire_stmt=branch(select(Tire),Tire);part_stmt=branch(select(Part),Part)
    occurrences=list(db.scalars(occurrence_stmt).all());tasks=list(db.scalars(task_stmt).all());maintenance=list(db.scalars(maintenance_stmt).all());tires=list(db.scalars(tire_stmt).all());parts=list(db.scalars(part_stmt).all())
    today=date.today()
    return {
        "period":{"start":start,"end":end},
        "operation":{"routes":len(routes),"completed":sum(1 for row in routes if row.status=="finalizada"),"in_transit":sum(1 for row in routes if row.status=="em_rota"),"cancelled":sum(1 for row in routes if row.status=="cancelada"),"completion_rate":round(100*sum(1 for row in routes if row.status=="finalizada")/len(routes),1) if routes else 0,"by_status":[{"label":key,"value":value} for key,value in route_status.items()]},
        "finance":{"revenue":revenue_total,"expense":expense_total,"result":round(revenue_total-expense_total,2),"margin":round(100*(revenue_total-expense_total)/revenue_total,1) if revenue_total else 0,"payable":round(pending_payable,2),"receivable":round(pending_receivable,2),"expenses_by_category":sum_by(expenses,lambda row:row.reason,lambda row:row.amount),"expenses_by_vehicle":sum_by(expenses,lambda row:vehicles.get(row.vehicle_id,"Sem veículo"),lambda row:row.amount)},
        "control":{"occurrences_open":sum(1 for row in occurrences if row.status not in {"finalizada","resolvida"}),"tasks_open":sum(1 for row in tasks if row.status=="open"),"tasks_in_progress":sum(1 for row in tasks if row.status=="in_progress"),"tasks_returned":sum(1 for row in tasks if row.status=="returned")},
        "fleet":{"maintenance_open":sum(1 for row in maintenance if row.status in {"aberta","em_andamento"}),"maintenance_overdue":sum(1 for row in maintenance if row.status not in {"concluida","cancelada"} and row.expected_completion_date and row.expected_completion_date<today),"tires_total":len(tires),"tires_attention":sum(1 for row in tires if row.status!="descartado" and row.tread_depth_mm is not None and row.tread_depth_mm<=3),"parts_low":sum(1 for row in parts if row.quantity<=row.minimum_quantity),"stock_value":round(sum(float(row.quantity or 0)*float(row.average_cost or 0) for row in parts),2)},
    }


class ProofReportOut(BaseModel):
    route_id: int
    codigo_ut: str
    route_date: date
    stop_id: int
    sequence: int
    customer_name: str
    status: str
    proof_type: str
    filename: str | None = None
    url: str | None = None


class ExpenseReportOut(BaseModel):
    id: int
    expense_date: date
    driver_name: str | None = None
    vehicle_plate: str | None = None
    reason: str
    amount: float | None = None
    route_id: int | None = None
    notes: str | None = None
    filename: str | None = None
    url: str | None = None


class AuditReportOut(BaseModel):
    id: int
    user_id: int | None = None
    user_name: str | None = None
    user_email: str | None = None
    action: str
    entity: str
    entity_id: str | None = None
    detail: str | None = None
    ip: str | None = None
    created_at: datetime


def _safe_filename(filename: str | None) -> str:
    name = Path(filename or "arquivo").name
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name)[:120] or "arquivo"


def _filename(attachment: Attachment | None) -> str | None:
    return Path(attachment.storage_key).name.split("-", 1)[-1] if attachment else None


def _proof_rows(db: Session, user: User, start: date | None, end: date | None) -> list[tuple[Route, RouteStop, str, Attachment]]:
    stmt = (
        select(Route)
        .options(
            selectinload(Route.stops).selectinload(RouteStop.proof_attachment),
            selectinload(Route.stops).selectinload(RouteStop.warehouse_return_attachment),
        )
        .order_by(Route.route_date.desc(), Route.id.desc())
    )
    stmt = scope_by_branch(stmt, Route.branch_id, user, db)
    if start:
        stmt = stmt.where(Route.route_date >= start)
    if end:
        stmt = stmt.where(Route.route_date <= end)
    rows: list[tuple[Route, RouteStop, str, Attachment]] = []
    for route in db.scalars(stmt).all():
        for stop in sorted(route.stops, key=lambda item: (item.sequence or 0, item.id)):
            if stop.proof_attachment:
                rows.append((route, stop, "entrega", stop.proof_attachment))
            if stop.warehouse_return_attachment:
                rows.append((route, stop, "devolucao_armazem", stop.warehouse_return_attachment))
    return rows


def _expense_rows(db: Session, user: User, start: date | None, end: date | None) -> list[Expense]:
    stmt = (
        select(Expense)
        .where(Expense.approval_status == "approved")
        .options(selectinload(Expense.driver), selectinload(Expense.attachment))
        .options(selectinload(Expense.vehicle))
        .order_by(Expense.expense_date.desc(), Expense.id.desc())
    )
    stmt = scope_by_branch(stmt, Expense.branch_id, user, db)
    if start:
        stmt = stmt.where(Expense.expense_date >= start)
    if end:
        stmt = stmt.where(Expense.expense_date <= end)
    return list(db.scalars(stmt).all())


def _audit_rows(db: Session, start: date | None, end: date | None) -> list[tuple[AuditLog, str | None, str | None]]:
    stmt = select(AuditLog, User.name, User.email).join(User, User.id == AuditLog.user_id, isouter=True)
    if start:
        stmt = stmt.where(AuditLog.created_at >= datetime.combine(start, time.min))
    if end:
        stmt = stmt.where(AuditLog.created_at <= datetime.combine(end, time.max))
    stmt = stmt.order_by(AuditLog.created_at.desc(), AuditLog.id.desc())
    return [(row.AuditLog, row.name, row.email) for row in db.execute(stmt).all()]


@router.get(
    "/routes.xlsx",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO, Role.AUDITOR))],
)
def export_routes_xlsx(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    stmt = (
        select(Route, RouteStop, Driver, Vehicle)
        .join(RouteStop, RouteStop.route_id == Route.id, isouter=True)
        .join(Driver, Driver.id == Route.driver_id, isouter=True)
        .join(Vehicle, Vehicle.id == Route.vehicle_id, isouter=True)
        .order_by(Route.route_date.desc(), Route.id.desc(), RouteStop.sequence)
    )
    stmt = scope_by_branch(stmt, Route.branch_id, user, db)
    if start:
        stmt = stmt.where(Route.route_date >= start)
    if end:
        stmt = stmt.where(Route.route_date <= end)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rotas"
    headers = [
        "Rota ID", "Codigo UT", "Data da rota", "Status", "Origem", "Morada origem",
        "Motorista", "Veiculo", "Entrega seq.", "Cliente", "Cidade", "Morada cliente",
        "Data planejada", "Hora planejada", "Status entrega", "Pedido", "Peso kg",
        "Paletes", "Tipo de parada", "Tipo devolução",
        "Quantidade devolvida", "Comprovante entrega", "Comprovante devolução armazém",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for route, stop, driver, vehicle in db.execute(stmt).all():
        sheet.append([
            route.id,
            route.codigo_ut,
            route.route_date.isoformat(),
            route.status,
            route.origin_name,
            route.origin_address,
            driver.name if driver else None,
            vehicle.plate if vehicle else None,
            stop.sequence if stop else None,
            stop.customer_name if stop else None,
            stop.city if stop else None,
            stop.customer_address if stop else None,
            stop.planned_date.isoformat() if stop and stop.planned_date else None,
            stop.planned_time.isoformat(timespec="minutes") if stop and stop.planned_time else None,
            stop.status if stop else None,
            stop.order_number if stop else None,
            stop.weight_kg if stop else None,
            stop.pallets if stop else None,
            stop.stop_type if stop else None,
            stop.return_type if stop else None,
            stop.returned_quantity if stop else None,
            _filename(stop.proof_attachment) if stop else None,
            _filename(stop.warehouse_return_attachment) if stop else None,
        ])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"relatorio-rotas-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/failures.xlsx",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.AUDITOR))],
)
def export_failures_xlsx(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    stmt = (
        select(Route, RouteStop, Driver, DeliveryFailureReason)
        .join(RouteStop, RouteStop.route_id == Route.id)
        .join(Driver, Driver.id == Route.driver_id, isouter=True)
        .join(DeliveryFailureReason, DeliveryFailureReason.id == RouteStop.failure_reason_id, isouter=True)
        .where(RouteStop.status == "falha")
        .order_by(Route.route_date.desc(), Route.id.desc(), RouteStop.sequence)
    )
    stmt = scope_by_branch(stmt, Route.branch_id, user, db)
    if start:
        stmt = stmt.where(Route.route_date >= start)
    if end:
        stmt = stmt.where(Route.route_date <= end)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Falhas"
    headers = [
        "Rota ID", "Codigo UT", "Data", "Motorista", "Parada", "Cliente", "Cidade",
        "Pedido", "Motivo", "Tipo devolução", "Quantidade devolvida", "Comprovante armazém",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for route, stop, driver, reason in db.execute(stmt).all():
        sheet.append([
            route.id,
            route.codigo_ut,
            route.route_date.isoformat(),
            driver.name if driver else None,
            stop.sequence,
            stop.customer_name,
            stop.city,
            stop.order_number,
            reason.label if reason else None,
            stop.return_type,
            stop.returned_quantity,
            _filename(stop.warehouse_return_attachment),
        ])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"relatorio-falhas-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get(
    "/proofs",
    response_model=list[ProofReportOut],
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.AUDITOR))],
)
def list_proofs(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return [
        ProofReportOut(
            route_id=route.id,
            codigo_ut=route.codigo_ut,
            route_date=route.route_date,
            stop_id=stop.id,
            sequence=stop.sequence,
            customer_name=stop.customer_name,
            status=stop.status,
            proof_type=proof_type,
            filename=_filename(attachment),
            url=storage.get_presigned_url(attachment.bucket, attachment.storage_key),
        )
        for route, stop, proof_type, attachment in _proof_rows(db, user, start, end)
    ]


@router.get(
    "/proofs.zip",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.AUDITOR))],
)
def download_proofs_zip(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for route, stop, proof_type, attachment in _proof_rows(db, user, start, end):
            filename = _safe_filename(_filename(attachment))
            archive_name = (
                f"{route.route_date.isoformat()}-{_safe_filename(route.codigo_ut)}/"
                f"parada-{stop.sequence}-{_safe_filename(stop.customer_name)}/"
                f"{proof_type}-{filename}"
            )
            archive.writestr(archive_name, storage.get_object_bytes(attachment.bucket, attachment.storage_key))
    output.seek(0)
    suffix = f"{start or 'inicio'}-{end or date.today().isoformat()}"
    return StreamingResponse(
        output,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="comprovantes-{suffix}.zip"'},
    )


@router.get(
    "/expenses",
    response_model=list[ExpenseReportOut],
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO, Role.AUDITOR))],
)
def list_expense_report(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    return [
        ExpenseReportOut(
            id=expense.id,
            expense_date=expense.expense_date,
            driver_name=expense.driver.name if expense.driver else None,
            vehicle_plate=expense.vehicle.plate if expense.vehicle else None,
            reason=expense.reason,
            amount=float(expense.amount) if expense.amount is not None else None,
            route_id=expense.route_id,
            notes=expense.notes,
            filename=_filename(expense.attachment),
            url=storage.get_presigned_url(expense.attachment.bucket, expense.attachment.storage_key) if expense.attachment else None,
        )
        for expense in _expense_rows(db, user, start, end)
    ]


@router.get(
    "/expenses.xlsx",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO, Role.AUDITOR))],
)
def export_expenses_report_xlsx(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Despesas"
    headers = ["ID", "Data", "Motorista", "Placa", "Motivo", "Valor", "Rota", "Observações", "Comprovante"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for expense in _expense_rows(db, user, start, end):
        sheet.append([
            expense.id,
            expense.expense_date.isoformat(),
            expense.driver.name if expense.driver else None,
            expense.vehicle.plate if expense.vehicle else None,
            expense.reason,
            float(expense.amount) if expense.amount is not None else None,
            expense.route_id,
            expense.notes,
            _filename(expense.attachment),
        ])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = f"{start or 'inicio'}-{end or date.today().isoformat()}"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="relatorio-despesas-{suffix}.xlsx"'},
    )


@router.get(
    "/expenses.zip",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO, Role.AUDITOR))],
)
def download_expense_proofs_zip(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        for expense in _expense_rows(db, user, start, end):
            attachment = expense.attachment
            if attachment is None:
                continue
            driver = _safe_filename(expense.driver.name if expense.driver else f"motorista-{expense.driver_id}")
            filename = _safe_filename(_filename(attachment))
            archive_name = f"{expense.expense_date.isoformat()}/{driver}/despesa-{expense.id}-{filename}"
            archive.writestr(archive_name, storage.get_object_bytes(attachment.bucket, attachment.storage_key))
    output.seek(0)
    suffix = f"{start or 'inicio'}-{end or date.today().isoformat()}"
    return StreamingResponse(
        output,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="comprovantes-despesas-{suffix}.zip"'},
    )


@router.get(
    "/audit",
    response_model=list[AuditReportOut],
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.AUDITOR))],
)
def list_audit_report(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
):
    return [
        AuditReportOut(
            id=log_row.id,
            user_id=log_row.user_id,
            user_name=name,
            user_email=email,
            action=log_row.action,
            entity=log_row.entity,
            entity_id=log_row.entity_id,
            detail=log_row.detail,
            ip=log_row.ip,
            created_at=log_row.created_at,
        )
        for log_row, name, email in _audit_rows(db, start, end)
    ]


@router.get(
    "/audit.xlsx",
    dependencies=[Depends(require_roles(Role.ADMIN_GLOBAL, Role.GESTOR_BRASIL, Role.AUDITOR))],
)
def export_audit_report_xlsx(
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Auditoria"
    headers = ["ID", "Data/hora", "Usuário", "E-mail", "Ação", "Entidade", "ID entidade", "Detalhe", "IP"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for log_row, name, email in _audit_rows(db, start, end):
        sheet.append([
            log_row.id,
            log_row.created_at.isoformat() if log_row.created_at else None,
            name,
            email,
            log_row.action,
            log_row.entity,
            log_row.entity_id,
            log_row.detail,
            log_row.ip,
        ])
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    suffix = f"{start or 'inicio'}-{end or date.today().isoformat()}"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="relatorio-auditoria-{suffix}.xlsx"'},
    )


class BalanceteRouteRow(BaseModel):
    route_id: int | None
    codigo_ut: str | None
    route_date: date | None
    revenue: float
    expense: float
    balance: float


class BalanceteOut(BaseModel):
    month: str
    revenue_total: float
    expense_total: float
    balance: float
    by_route: list[BalanceteRouteRow]


_FINANCEIRO = require_permission(FINANCE_VIEW, Role.GESTOR_BRASIL, Role.GESTOR_FINANCEIRO)


@router.get("/balancete", response_model=BalanceteOut, dependencies=[Depends(_FINANCEIRO)])
def balancete(month: str, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    """Receita x despesa do mês — total geral e detalhado por rota."""
    try:
        year, month_num = [int(part) for part in month.split("-", 1)]
    except ValueError:
        raise HTTPException(status_code=400, detail="Mês inválido. Use AAAA-MM.")

    revenue_stmt = select(Revenue).where(
        extract("year", Revenue.revenue_date) == year, extract("month", Revenue.revenue_date) == month_num,
    )
    expense_stmt = select(Expense).where(
        extract("year", Expense.expense_date) == year, extract("month", Expense.expense_date) == month_num,
        Expense.approval_status == "approved",
    )
    revenue_stmt = scope_by_branch(revenue_stmt, Revenue.branch_id, user, db)
    expense_stmt = scope_by_branch(expense_stmt, Expense.branch_id, user, db)

    revenues = db.scalars(revenue_stmt).all()
    expenses = db.scalars(expense_stmt).all()

    route_ids = {r.route_id for r in revenues if r.route_id} | {e.route_id for e in expenses if e.route_id}
    routes = {r.id: r for r in db.scalars(select(Route).where(Route.id.in_(route_ids))).all()} if route_ids else {}

    by_route: dict[int | None, dict] = {}
    for revenue_row in revenues:
        bucket = by_route.setdefault(revenue_row.route_id, {"revenue": 0.0, "expense": 0.0})
        bucket["revenue"] += float(revenue_row.amount)
    for expense_row in expenses:
        bucket = by_route.setdefault(expense_row.route_id, {"revenue": 0.0, "expense": 0.0})
        bucket["expense"] += float(expense_row.amount or 0)

    rows = [
        BalanceteRouteRow(
            route_id=route_id,
            codigo_ut=routes[route_id].codigo_ut if route_id in routes else None,
            route_date=routes[route_id].route_date if route_id in routes else None,
            revenue=round(values["revenue"], 2),
            expense=round(values["expense"], 2),
            balance=round(values["revenue"] - values["expense"], 2),
        )
        for route_id, values in sorted(by_route.items(), key=lambda kv: (kv[1]["revenue"] - kv[1]["expense"]))
    ]
    revenue_total = round(sum(r.revenue for r in rows), 2)
    expense_total = round(sum(r.expense for r in rows), 2)
    return BalanceteOut(
        month=month, revenue_total=revenue_total, expense_total=expense_total,
        balance=round(revenue_total - expense_total, 2), by_route=rows,
    )
