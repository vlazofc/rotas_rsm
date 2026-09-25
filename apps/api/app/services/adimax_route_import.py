"""Importador do modelo Gestão Adimax 2026 / aba LAST MILE - ADIMAX."""
from __future__ import annotations

import json
import re
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
import io
from sqlalchemy import select

from app.db.models import (Branch, Customer, DeliveryDestination, Driver, Route,
                           RouteOrigin, RouteStop, Vehicle)
from app.db.session import SessionLocal
from app.services.import_changes import snapshot_import, collect_changes, authorize_changes
from app.services.route_carrier import apply_import_carrier, resolve_import_carrier

SHEET = "LAST MILE - ADIMAX"
ADIMAX_HEADERS = ["DATA CARREGAMENTO", "DATA ENTREGA", "ROTA", "MOTORISTA", "CPF", "PLACA",
    "TIPO MOTORISTA", "PERFIL ENVIADO", "PERFIL SOLICITADO", "VISÃO TIPOLOGIA", "PERNOITE",
    "DIARIA", "OBSERVAÇÃO", "AJUDANTE SOLICITADO", "AJUDATE ENVIADO", "QUANTIDADE",
    "SEQUENCIA", "CARGA", "PEDIDO", "NOTA FISCAL", "CTE", "ORIGEM", "CLIENTE", "DESTINO",
    "NUMERO", "BAIRRO", "CIDADE", "PESO BRUTO", "OBSERVAÇÃO CLIENTE", "OBSERVAÇÃO CLIENTE 2",
    "ID TRANSPORTADORA", "TRANSPORTADORA"]


def build_adimax_template_xlsx() -> io.BytesIO:
    workbook = openpyxl.Workbook()
    sheet = workbook.active; sheet.title = SHEET; sheet.append(ADIMAX_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="F9A61A")
        sheet.column_dimensions[cell.column_letter].width = max(14, min(28, len(str(cell.value)) + 3))
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = f"A1:AD1"
    guide = workbook.create_sheet("INSTRUÇÕES")
    instructions = [
        ("CARGA", "Código único da rota. Obrigatório."),
        ("SEQUENCIA", "Ordem original das paradas definida pelo cliente. Obrigatório."),
        ("DATA", "Preencher no formato dd/mm/aaaa; operação no fuso de Brasília."),
        ("ENDEREÇO", "DESTINO + NUMERO + BAIRRO + CIDADE formam o endereço da parada."),
        ("ORIGEM", "Coluna V: endereço completo do CD de carregamento, com cidade e UF. Ponto inicial antes da entrega 1; não altera a sequência das entregas."),
        ("ADMINISTRATIVO", "PERNOITE, DIARIA, NOTA FISCAL e CTE são preenchidos por administradores/gestores."),
    ]
    guide.append(["CAMPO", "ORIENTAÇÃO"])
    for row in instructions: guide.append(row)
    for cell in guide[1]: cell.font = Font(bold=True)
    guide.column_dimensions["A"].width = 22; guide.column_dimensions["B"].width = 95
    output = io.BytesIO(); workbook.save(output); output.seek(0); return output


def _text(value) -> str | None:
    if value is None or str(value).strip() == "": return None
    if isinstance(value, float) and value.is_integer(): return str(int(value))
    return str(value).strip()


def _date(value) -> date | None:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    if not value: return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try: return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError: pass
    return None


def _number(value) -> float | None:
    if value in (None, ""): return None
    try: return float(str(value).replace(".", "").replace(",", ".") if "," in str(value) else value)
    except (TypeError, ValueError): return None


def _bool(value) -> bool | None:
    text = (_text(value) or "").upper()
    if not text: return None
    return text in {"SIM", "S", "1", "TRUE", "VERDADEIRO", "OK"}


def _plate(value) -> str | None:
    text = re.sub(r"[^A-Z0-9]", "", (_text(value) or "").upper())
    return text if len(text) >= 7 else None


def _document(value) -> str | None:
    text = re.sub(r"\D", "", _text(value) or "")
    return text or None


def is_adimax_workbook(source) -> bool:
    source.seek(0)
    workbook = openpyxl.load_workbook(source, read_only=True)
    found = SHEET in workbook.sheetnames
    workbook.close(); source.seek(0)
    return found


def import_adimax_routes(source, branch_id: int, *, confirmed: bool = False, user_id: int | None = None) -> dict:
    source.seek(0)
    sheet = openpyxl.load_workbook(source, read_only=True, data_only=True)[SHEET]
    headers = [_text(cell.value) or f"COL_{index}" for index, cell in enumerate(sheet[1], start=1)]
    rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows(min_row=2)
            if any(cell.value not in (None, "") for cell in row)]
    stats = {"rows_read": len(rows), "routes_created": 0, "routes_updated": 0,
             "stops_created": 0, "stops_updated": 0, "routes_optimized": 0,
             "routing_errors": 0, "route_ids": [], "drivers_created": 0,
             "vehicles_created": 0, "customers_created": 0, "destinations_created": 0,
             "origins_created": 0, "carrier_pending": 0, "errors": []}
    groups: dict[str, list[tuple]] = {}
    for index, row in enumerate(rows, start=2):
        load = _text(row[17]) if len(row) > 17 else None
        if not load:
            stats["errors"].append(f"Linha {index}: CARGA vazia; registro preservado apenas na planilha.")
            continue
        groups.setdefault(load, []).append(row)

    with SessionLocal() as db:
        changes = []
        branch = db.get(Branch, branch_id)
        if branch is None or branch.tenant_id is None: raise RuntimeError("Filial/cliente não configurado.")
        tenant_id = branch.tenant_id
        for load, group in groups.items():
            first = group[0]
            normalized_headers = [re.sub(r"\s+", " ", re.sub(r"[^A-Z0-9 ]", "", header.upper())).strip() for header in headers]
            carrier_id_index = next((i for i, header in enumerate(normalized_headers) if header == "ID TRANSPORTADORA"), None)
            carrier_name_index = next((i for i, header in enumerate(normalized_headers) if header == "TRANSPORTADORA"), None)
            route = db.scalar(select(Route).where(Route.branch_id == branch_id, Route.codigo_ut == load))
            new_route = route is None
            before_route = snapshot_import(route)
            if new_route:
                route = Route(branch_id=branch_id, tenant_id=tenant_id, codigo_ut=load,
                              route_date=_date(first[0]) or date.today(), source="automatico")
                db.add(route)
            route.route_date = _date(first[0]) or route.route_date
            route.delivery_date = next((_date(row[1]) for row in group if _date(row[1])), None)
            route.spreadsheet_route = _text(first[2])
            route.driver_type = _text(first[6]); route.vehicle_profile_sent = _text(first[7])
            route.vehicle_profile_requested = _text(first[8]); route.typology_view = _text(first[9])
            route.overnight = _bool(first[10]); route.daily_count = _number(first[11]); route.daily_value = _text(first[11])
            route.administrative_notes = _text(first[12]); route.helper_requested = _text(first[13])
            route.helper_sent = _text(first[14]); route.load_quantity = int(_number(first[15]) or 0) or None
            if new_route and route.delivery_date and route.delivery_date < date.today(): route.status = "finalizada"
            # A planilha pode vincular cadastros já existentes somente ao criar
            # uma rota. Reimportações nunca trocam escala nem alteram/criam
            # motoristas ou veículos mantidos pela operação.
            driver_name, cpf = _text(first[3]), _document(first[4])
            if new_route and driver_name:
                driver = db.scalar(select(Driver).where(Driver.branch_id == branch_id, Driver.document == cpf)) if cpf else None
                driver = driver or db.scalar(select(Driver).where(Driver.branch_id == branch_id, Driver.name.ilike(driver_name)))
                if driver is not None: route.driver_id = driver.id
            plate = _plate(first[5])
            if new_route and plate:
                vehicle = db.scalar(select(Vehicle).where(Vehicle.branch_id == branch_id, Vehicle.plate == plate))
                if vehicle is not None: route.vehicle_id = vehicle.id

            # O CD da primeira parada na sequência é o início da carga.
            ordered_rows = sorted(group, key=lambda row: _number(row[16]) or float('inf'))
            origin_address = next((_text(row[21]) for row in ordered_rows if len(row) > 21 and _text(row[21])), None)
            if origin_address:
                origin = db.scalar(select(RouteOrigin).where(RouteOrigin.tenant_id == tenant_id, RouteOrigin.address == origin_address))
                if origin is None:
                    origin = RouteOrigin(tenant_id=tenant_id, name=f"CD {origin_address.split(',')[0]}", address=origin_address)
                    db.add(origin); db.flush(); stats["origins_created"] += 1
                route.origin_id, route.origin_name, route.origin_address = origin.id, origin.name, origin.address
            collect_changes(changes, before_route, route, load)
            carrier, carrier_issue = resolve_import_carrier(
                db, branch_id,
                carrier_id=int(_number(first[carrier_id_index])) if carrier_id_index is not None and len(first) > carrier_id_index and _number(first[carrier_id_index]) is not None else None,
                name=_text(first[carrier_name_index]) if carrier_name_index is not None and len(first) > carrier_name_index else None,
            )
            apply_import_carrier(route, carrier, carrier_issue)
            if carrier_issue:
                stats["carrier_pending"] += 1
                stats["errors"].append(f"Rota {load}: {carrier_issue}")
            db.flush(); stats["route_ids"].append(route.id)
            stats["routes_created" if new_route else "routes_updated"] += 1

            existing = {stop.order_number: stop for stop in route.stops if stop.order_number}
            for position, row in enumerate(group, start=1):
                order = _text(row[18]); customer_name = _text(row[22]) or "Cliente não informado"
                customer = db.scalar(select(Customer).where(Customer.tenant_id == tenant_id, Customer.name.ilike(customer_name)))
                if customer is None:
                    customer = Customer(tenant_id=tenant_id, name=customer_name)
                    db.add(customer); db.flush(); stats["customers_created"] += 1
                street, number, district, city = (_text(row[23]), _text(row[24]), _text(row[25]), _text(row[26]))
                street = street or "Endereço não informado"
                destination = db.scalar(select(DeliveryDestination).where(
                    DeliveryDestination.customer_id == customer.id, DeliveryDestination.address == street,
                    DeliveryDestination.number == number, DeliveryDestination.district == district,
                    DeliveryDestination.city == city))
                if destination is None:
                    destination = DeliveryDestination(tenant_id=tenant_id, customer_id=customer.id,
                        label=_text(row[23]), address=street, number=number, district=district, city=city)
                    db.add(destination); db.flush(); stats["destinations_created"] += 1
                sequence = int(_number(row[16]) or position)
                stop = existing.get(order) if order else None
                new_stop = stop is None
                before_stop = snapshot_import(stop)
                if new_stop: stop = RouteStop(route_id=route.id, customer_name=customer_name)
                stop.sequence = sequence; stop.spreadsheet_sequence = sequence
                stop.customer_id = customer.id; stop.destination_id = destination.id
                stop.customer_name = customer_name; stop.client_name = customer_name
                stop.customer_address = ", ".join(part for part in (street, number, district) if part)
                stop.city = city; stop.order_number = order; stop.invoice_number = _text(row[19])
                stop.cte_number = _text(row[20]); stop.weight_kg = _number(row[27])
                stop.customer_notes = _text(row[28]); stop.customer_notes_2 = _text(row[29])
                stop.customer_notes_3 = _text(row[30]) if len(row) > 30 else None
                stop.raw_import_json = json.dumps(dict(zip(headers, (_text(value) for value in row))), ensure_ascii=False)
                if new_stop and route.status == "finalizada": stop.status = "entregue"
                if new_stop: db.add(stop)
                collect_changes(changes, before_stop, stop, load)
                if order: existing[order] = stop
                stats["stops_created" if new_stop else "stops_updated"] += 1
            db.flush()
        authorize_changes(db, changes, confirmed=confirmed, user_id=user_id)
        db.commit()
    return stats
