"""Importador da planilha de cargas JM (Planilha1 + Planilha2)."""
from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime

import openpyxl
import io
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.db.models import Branch, Route, RouteStop
from app.db.session import SessionLocal
from app.services.import_changes import authorize_changes, collect_changes, snapshot_import


def _text(value) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value) -> str:
    text = unicodedata.normalize("NFD", _text(value) or "")
    return re.sub(r"\s+", " ", "".join(c for c in text if unicodedata.category(c) != "Mn").upper()).strip()


def _number(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        text = str(value).strip()
        return float(text.replace(".", "").replace(",", ".") if "," in text else text)
    except (TypeError, ValueError):
        return None


def _date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    for fmt in ("%d/%m/%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def _canonical_headers(values) -> list[str]:
    headers = [_header(value) for value in values]
    if "NUMERO PEDIDO" not in headers:
        number_indexes = [index for index, name in enumerate(headers) if name == "NUMERO"]
        # A coluna de pedido fica no bloco inicial, antes de NF/cliente. Já o
        # número do endereço aparece imediatamente depois de LOGRADOURO.
        order_index = next((index for index in number_indexes
                            if "LOGRADOURO" not in headers[:index]), None)
        if order_index is not None:
            headers[order_index] = "NUMERO PEDIDO"
    return headers


def _rows(sheet) -> list[dict]:
    iterator = sheet.iter_rows(values_only=True)
    headers = _canonical_headers(next(iterator, ()))
    return [dict(zip(headers, row)) for row in iterator if any(value not in (None, "") for value in row)]


def build_jm_template_xlsx() -> io.BytesIO:
    """Gera o modelo no mesmo formato da carga operacional enviada pela JM."""
    workbook = openpyxl.Workbook()
    detail = workbook.active
    detail.title = "Detalhes da carga"
    detail.append([
        "Sequência", "Nº Carga", "Número pedido", "NF", "Razão Social / Nome",
        "Logradouro", "Número", "Bairro", "Cidade", "UF", "PESO BRUTO",
        "Observação de Entrega", "Observação Representante",
    ])
    fiscal = workbook.create_sheet("Notas e datas")
    fiscal.append([
        "Razão Social / Nome", "Nº Lote", "Número do pedido", "Nº Carga",
        "Data Emissão NFe", "Nº NF", "Situação NFe", "Representante Comercial",
        "Data Emissão Pedido", "Data Prev. Entrega",
    ])
    instructions = workbook.create_sheet("Instruções")
    instructions.append(["ORIENTAÇÃO", "DETALHE"])
    instructions.append(["Nome das abas", "Pode ser alterado; a identificação é feita pelas colunas."])
    instructions.append(["Número do pedido", "Também pode ser informado como Número ou Número pedido."])
    instructions.append(["Detalhes da carga", "Uma linha por entrega. Sequência, Nº Carga e pedido são obrigatórios."])
    instructions.append(["Notas e datas", "Use o mesmo pedido da primeira aba para atualizar NF e datas, inclusive retroativamente."])
    for sheet in (detail, fiscal, instructions):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F9A61A")
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = max(16, min(36, len(str(cell.value)) + 4))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    instructions.column_dimensions["B"].width = 95
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def is_jm_load_workbook(source) -> bool:
    source.seek(0)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            headers = set(_canonical_headers(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1), ())))
            if {"SEQUENCIA", "Nº CARGA", "NUMERO PEDIDO", "RAZAO SOCIAL / NOME", "LOGRADOURO", "CIDADE"} <= headers:
                return True
        return False
    finally:
        workbook.close()
        source.seek(0)


def import_jm_loads(source, branch_id: int, *, confirmed: bool = False, user_id: int | None = None) -> dict:
    source.seek(0)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    sheets = {sheet.title: _rows(sheet) for sheet in workbook.worksheets}
    workbook.close()
    detail_rows = next((rows for rows in sheets.values() if rows and "Nº CARGA" in rows[0] and "LOGRADOURO" in rows[0]), [])
    metadata_rows = next((rows for rows in sheets.values() if rows and "Nº CARGA" in rows[0] and "DATA PREV. ENTREGA" in rows[0]), [])
    metadata = {
        _text(row.get("N° DO PEDIDO") or row.get("Nº DO PEDIDO") or row.get("NUMERO DO PEDIDO") or row.get("NUMERO PEDIDO")): row
        for row in metadata_rows
    }
    stats = {"rows_read": len(detail_rows), "routes_created": 0, "routes_updated": 0,
             "stops_created": 0, "stops_updated": 0, "routes_optimized": 0,
             "routing_errors": 0, "route_ids": [], "errors": []}
    groups: dict[str, list[dict]] = {}
    for index, row in enumerate(detail_rows, 2):
        load = _text(row.get("Nº CARGA"))
        order = _text(row.get("NUMERO PEDIDO"))
        if not load or not order:
            stats["errors"].append(f"Linha {index}: Nº Carga ou Número pedido vazio — ignorada.")
            continue
        groups.setdefault(load, []).append(row)

    with SessionLocal() as db:
        branch = db.get(Branch, branch_id)
        if branch is None:
            raise RuntimeError("Filial não encontrada.")
        changes = []
        for load, rows in groups.items():
            first_meta = next((metadata.get(_text(row.get("NUMERO PEDIDO"))) for row in rows if metadata.get(_text(row.get("NUMERO PEDIDO")))), {})
            planned_date = _date(first_meta.get("DATA PREV. ENTREGA")) or _date(first_meta.get("DATA EMISSAO NFE")) or date.today()
            route = db.scalar(select(Route).where(Route.branch_id == branch_id, Route.codigo_ut == load))
            is_new_route = route is None
            before_route = snapshot_import(route)
            if is_new_route:
                route = Route(branch_id=branch_id, tenant_id=branch.tenant_id, codigo_ut=load,
                              route_date=planned_date, delivery_date=planned_date,
                              spreadsheet_route=load, source="automatico", status="planejada")
                db.add(route)
            else:
                route.route_date = planned_date
                route.delivery_date = planned_date
            collect_changes(changes, before_route, route, load)
            db.flush()
            stats["route_ids"].append(route.id)
            stats["routes_created" if is_new_route else "routes_updated"] += 1
            existing = {stop.order_number: stop for stop in route.stops if stop.order_number}
            for position, row in enumerate(sorted(rows, key=lambda item: _number(item.get("SEQUENCIA")) or 999999), 1):
                order = _text(row.get("NUMERO PEDIDO"))
                meta = metadata.get(order, {})
                stop = existing.get(order)
                is_new_stop = stop is None
                before_stop = snapshot_import(stop)
                if is_new_stop:
                    stop = RouteStop(route_id=route.id, customer_name="Cliente não informado")
                    db.add(stop)
                sequence = int(_number(row.get("SEQUENCIA")) or position)
                street, number, district = _text(row.get("LOGRADOURO")), _text(row.get("NUMERO")), _text(row.get("BAIRRO"))
                stop.sequence = stop.spreadsheet_sequence = sequence
                stop.order_number = order
                stop.invoice_number = _text(row.get("NF") or meta.get("Nº NF"))
                stop.customer_name = stop.client_name = _text(row.get("RAZAO SOCIAL / NOME")) or stop.customer_name
                stop.customer_address = ", ".join(value for value in (street, number, district) if value)
                stop.city = _text(row.get("CIDADE"))
                stop.province = _text(row.get("UF"))
                stop.weight_kg = _number(row.get("PESO BRUTO"))
                stop.planned_date = _date(meta.get("DATA PREV. ENTREGA")) or planned_date
                stop.invoicing_date = _date(meta.get("DATA EMISSAO NFE"))
                stop.customer_notes = _text(row.get("OBSERVACAO DE ENTREGA"))
                stop.customer_notes_2 = _text(row.get("OBSERVACAO REPRESENTANTE"))
                raw = {**{key: _text(value) for key, value in row.items()},
                       **{f"Planilha2 - {key}": _text(value) for key, value in meta.items()}}
                stop.raw_import_json = json.dumps(raw, ensure_ascii=False)
                collect_changes(changes, before_stop, stop, load)
                db.flush()
                stats["stops_created" if is_new_stop else "stops_updated"] += 1
        authorize_changes(db, changes, confirmed=confirmed, user_id=user_id)
        db.commit()
    return stats
