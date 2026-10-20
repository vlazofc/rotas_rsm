import io

import openpyxl
import pytest
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.models import Branch, Carrier, CarrierBranch, Route, Tenant
from app.db.session import Base
from app.services import jm_load_import
from app.services.jm_load_import import build_jm_template_xlsx
from app.services.route_carrier import resolve_import_carrier


@pytest.fixture
def carrier_identity_db():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    with Session(engine) as db:
        tenant = Tenant(name="Adimax", slug="adimax-import-identity")
        db.add(tenant); db.flush()
        salto = Branch(name="Salto", tenant_id=tenant.id)
        barueri = Branch(name="Barueri", tenant_id=tenant.id)
        db.add_all([salto, barueri]); db.flush()
        alfa = Carrier(name="Transportadora Alfa", document="11111111000111", person_type="pessoa_juridica", tenant_id=tenant.id)
        db.add(alfa); db.flush()
        db.add(CarrierBranch(carrier_id=alfa.id, branch_id=salto.id))
        db.commit()
        yield db, salto, barueri, alfa
    engine.dispose()


def test_import_resolves_carrier_by_id_and_checks_name(carrier_identity_db):
    db, salto, _, alfa = carrier_identity_db
    carrier, issue = resolve_import_carrier(db, salto.id, carrier_id=alfa.id, name="Transportadora Alfa")
    assert carrier.id == alfa.id
    assert issue is None

    carrier, issue = resolve_import_carrier(db, salto.id, carrier_id=alfa.id, name="Transportadora Beta")
    assert carrier is None
    assert "não corresponde" in issue


def test_import_rejects_unknown_id_and_carrier_outside_branch(carrier_identity_db):
    db, salto, barueri, alfa = carrier_identity_db
    carrier, issue = resolve_import_carrier(db, salto.id, carrier_id=999999, name="Transportadora Alfa")
    assert carrier is None
    assert "inexistente ou inativa" in issue

    carrier, issue = resolve_import_carrier(db, barueri.id, carrier_id=alfa.id, name="Transportadora Alfa")
    assert carrier is None
    assert "não habilitada" in issue


def test_template_lists_carrier_ids_branch_names_and_instructions():
    workbook = openpyxl.load_workbook(build_jm_template_xlsx(
        [(7, "Transportadora Alfa"), (9, "Transportadora Beta")],
        [(10, "Barueri"), (20, "Salto")],
    ))
    try:
        detail_headers = [cell.value for cell in workbook["Detalhes da carga"][1]]
        assert detail_headers == [
            "Transportadora", "Filial", "Sequência", "Nº Carga", "Número pedido", "NF",
            "Razão Social / Nome", "Logradouro", "Número", "Bairro", "Cidade", "UF",
            "PESO BRUTO", "ID Transportadora", "Observação de Entrega", "Observação Representante",
        ]
        assert list(workbook["Transportadoras"].values) == [
            ("ID TRANSPORTADORA", "TRANSPORTADORA"),
            (7, "Transportadora Alfa"),
            (9, "Transportadora Beta"),
        ]
        assert list(workbook["Filiais"].values) == [
            ("ID FILIAL", "FILIAL"),
            (10, "Barueri"),
            (20, "Salto"),
        ]
        instructions = dict(workbook["Instruções"].values)
        assert "direciona automaticamente" in instructions["Filial"]
    finally:
        workbook.close()


def test_same_template_layout_serves_adimax_and_carrier_scope():
    adimax = openpyxl.load_workbook(build_jm_template_xlsx([(7, "Alfa"), (9, "Beta")]))
    carrier = openpyxl.load_workbook(build_jm_template_xlsx([(7, "Alfa")]))
    try:
        assert adimax.sheetnames == carrier.sheetnames == [
            "Detalhes da carga", "Notas e datas", "Instruções", "Filiais", "Transportadoras",
        ]
        assert list(adimax["Detalhes da carga"].values) == list(carrier["Detalhes da carga"].values)
        assert list(adimax["Transportadoras"].values)[1:] == [(7, "Alfa"), (9, "Beta")]
        assert list(carrier["Transportadoras"].values)[1:] == [(7, "Alfa")]
    finally:
        adimax.close()
        carrier.close()


def test_jm_import_routes_each_load_to_branch_named_in_spreadsheet(carrier_identity_db, monkeypatch):
    db, salto, barueri, alfa = carrier_identity_db
    engine = db.get_bind()
    workbook = openpyxl.load_workbook(build_jm_template_xlsx(
        [(alfa.id, alfa.name)],
        [(salto.id, salto.name), (barueri.id, barueri.name)],
    ))
    detail = workbook["Detalhes da carga"]
    detail.append([
        alfa.name, "Salto", 1, "CARGA-001", "PEDIDO-001", "NF-001",
        "Cliente Teste", "Rua Um", "10", "Centro", "Sorocaba", "SP",
        100, alfa.id, None, None,
    ])
    source = io.BytesIO()
    workbook.save(source)
    workbook.close()
    source.seek(0)

    monkeypatch.setattr(jm_load_import, "SessionLocal", lambda: Session(engine))
    result = jm_load_import.import_jm_loads(
        source,
        allowed_branch_ids={salto.id, barueri.id},
        confirmed=True,
    )

    with Session(engine) as verification:
        route = verification.scalar(select(Route).where(Route.codigo_ut == "CARGA-001"))
        assert route is not None
        assert route.branch_id == salto.id
    assert result["routes_created"] == 1


def test_jm_import_rejects_unknown_branch(carrier_identity_db, monkeypatch):
    db, salto, barueri, alfa = carrier_identity_db
    engine = db.get_bind()
    workbook = openpyxl.load_workbook(build_jm_template_xlsx())
    workbook["Detalhes da carga"].append([
        alfa.name, "Filial Inexistente", 1, "CARGA-002", "PEDIDO-002", None,
        "Cliente Teste", "Rua Dois", "20", "Centro", "Sorocaba", "SP",
        50, alfa.id, None, None,
    ])
    source = io.BytesIO()
    workbook.save(source)
    workbook.close()
    source.seek(0)

    monkeypatch.setattr(jm_load_import, "SessionLocal", lambda: Session(engine))
    with pytest.raises(RuntimeError, match="não encontrada, inativa ou fora do acesso"):
        jm_load_import.import_jm_loads(
            source,
            allowed_branch_ids={salto.id, barueri.id},
            confirmed=True,
        )
