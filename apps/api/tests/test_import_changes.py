import io
import json
from datetime import date

import pytest
import openpyxl
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.session import Base
from app.db.models import AuditLog, Branch, Driver, Route, RouteStop, Tenant, User, Vehicle
from app.services import generic_route_import as importer
from app.services import adimax_route_import as adimax
from app.services import jm_load_import as jm_load
from app.services.import_changes import ImportConfirmationRequired


@pytest.fixture
def context(monkeypatch):
    engine = create_engine("sqlite://", poolclass=StaticPool)
    Base.metadata.create_all(engine)
    monkeypatch.setattr(importer, "SessionLocal", lambda: Session(engine))
    monkeypatch.setattr(adimax, "SessionLocal", lambda: Session(engine))
    monkeypatch.setattr(jm_load, "SessionLocal", lambda: Session(engine))
    with Session(engine) as db:
        tenant = Tenant(name="Test", slug="test")
        db.add(tenant); db.flush()
        branch = Branch(name="Test", tenant_id=tenant.id)
        user = User(name="Autorizador", email="import@example.com", role="admin_global")
        db.add_all([branch, user]); db.flush()
        route = Route(branch_id=branch.id, codigo_ut="123", route_date=date(2026, 9, 10), status="em_rota")
        db.add(route); db.flush()
        db.add(RouteStop(route_id=route.id, sequence=1, order_number="456", customer_name="Cliente", city="Antiga", status="entregue"))
        db.commit()
        yield engine, branch.id, user.id
    engine.dispose()


def upload(branch_id, **kwargs):
    source = io.BytesIO("DATA ROTA,ROTA,PEDIDO,DESTINO,CIDADE\n10/09/2026,123,456,Cliente,Nova\n".encode())
    return importer.import_generic_routes(source, "test.csv", branch_id, **kwargs)


def test_confirmation_rolls_back(context):
    engine, branch_id, _ = context
    with pytest.raises(ImportConfirmationRequired) as exc:
        upload(branch_id)
    assert any(c["field"] == "city" and c["old"] == "Antiga" and c["new"] == "Nova" for c in exc.value.changes)
    with Session(engine) as db:
        assert db.scalar(select(RouteStop)).city == "Antiga"
        assert not db.scalars(select(AuditLog)).all()


def test_confirmed_import_audits_and_preserves_progress(context):
    engine, branch_id, user_id = context
    upload(branch_id, confirmed=True, user_id=user_id)
    with Session(engine) as db:
        stop = db.scalar(select(RouteStop))
        assert stop.city == "Nova"
        assert stop.status == "entregue"
        assert db.scalar(select(Route)).status == "em_rota"
        audit = db.scalar(select(AuditLog).where(AuditLog.action == "import_overwrite"))
        assert audit.user_id == user_id
        assert json.loads(audit.detail)["old"] == "Antiga"
    # Identical reimport must not require authorization or duplicate the stop.
    upload(branch_id)
    with Session(engine) as db:
        assert len(db.scalars(select(RouteStop)).all()) == 1


def test_confirmation_requires_identified_user(context):
    _, branch_id, _ = context
    with pytest.raises(ImportConfirmationRequired):
        upload(branch_id, confirmed=True)


def test_reimport_preserves_assignment_and_master_data(context):
    engine, branch_id, user_id = context
    with Session(engine) as db:
        driver = Driver(branch_id=branch_id, name="Motorista Operacional")
        vehicle = Vehicle(branch_id=branch_id, plate="ABC1D23")
        db.add_all([driver, vehicle]); db.flush()
        route = db.scalar(select(Route))
        route.driver_id, route.vehicle_id = driver.id, vehicle.id
        db.commit()
        expected = route.driver_id, route.vehicle_id
    source = io.BytesIO("DATA ROTA,ROTA,PEDIDO,DESTINO,CIDADE,MOTORISTA,PLACA\n10/09/2026,123,456,Cliente,Nova,Novo Motorista,ZZZ9Z99\n".encode())
    importer.import_generic_routes(source, "test.csv", branch_id, confirmed=True, user_id=user_id)
    with Session(engine) as db:
        route = db.scalar(select(Route))
        assert (route.driver_id, route.vehicle_id) == expected
        assert [row.name for row in db.scalars(select(Driver)).all()] == ["Motorista Operacional"]
        assert [row.plate for row in db.scalars(select(Vehicle)).all()] == ["ABC1D23"]


def test_adimax_overwrite_confirmation_and_idempotency(context):
    engine, branch_id, user_id = context
    with Session(engine) as db:
        driver = Driver(branch_id=branch_id, name="Motorista Operacional")
        vehicle = Vehicle(branch_id=branch_id, plate="ABC1D23")
        db.add_all([driver, vehicle]); db.flush()
        route = db.scalar(select(Route))
        route.driver_id, route.vehicle_id = driver.id, vehicle.id
        db.commit()
        expected_assignment = route.driver_id, route.vehicle_id
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = adimax.SHEET
    sheet.append(adimax.ADIMAX_HEADERS)
    row = [None] * 30
    row[0], row[1], row[16], row[17], row[18] = date(2026, 9, 10), date(2020, 1, 1), 1, "123", "456"
    row[3], row[4], row[5] = "Novo Motorista", "12345678900", "ZZZ9Z99"
    row[22], row[23], row[26] = "Cliente", "Rua Teste", "Nova"
    sheet.append(row)
    source = io.BytesIO()
    workbook.save(source)
    with pytest.raises(ImportConfirmationRequired):
        adimax.import_adimax_routes(source, branch_id)
    with Session(engine) as db:
        assert db.scalar(select(RouteStop)).city == "Antiga"
    adimax.import_adimax_routes(source, branch_id, confirmed=True, user_id=user_id)
    adimax.import_adimax_routes(source, branch_id)
    with Session(engine) as db:
        assert db.scalar(select(Route)).status == "em_rota"
        assert (db.scalar(select(Route)).driver_id, db.scalar(select(Route)).vehicle_id) == expected_assignment
        assert db.scalar(select(RouteStop)).status == "entregue"
        assert len(db.scalars(select(RouteStop)).all()) == 1
        assert len(db.scalars(select(Driver)).all()) == 1
        assert len(db.scalars(select(Vehicle)).all()) == 1
        assert db.scalar(select(AuditLog)).user_id == user_id


def test_jm_load_accepts_numero_as_order_and_keeps_address_number(context):
    engine, branch_id, _ = context
    workbook = openpyxl.Workbook()
    detail = workbook.active
    detail.title = "Planilha1"
    detail.append(["Sequência", "Nº Carga", "Número", "NF", "Razão Social / Nome",
                   "Logradouro", "Número", "Bairro", "Cidade", "UF", "PESO BRUTO",
                   "Observação de Entrega", "Observação Representante"])
    detail.append([1, "542470", "14397867", "851557", "Cliente JM", "Rua Teste", "26",
                   "Centro", "Piedade", "SP", 39.1, "Entregar cedo", "Representante"])
    metadata = workbook.create_sheet("Planilha2")
    metadata.append(["Razão Social / Nome", "Número", "Nº Carga", "Data Emissão NFe", "Data Prev. Entrega"])
    metadata.append(["Cliente JM", "14397867", "542470", date(2026, 9, 10), date(2026, 9, 14)])
    source = io.BytesIO(); workbook.save(source)
    assert jm_load.is_jm_load_workbook(source)
    result = jm_load.import_jm_loads(source, branch_id)
    assert result["routes_created"] == result["stops_created"] == 1
    with Session(engine) as db:
        route = db.scalar(select(Route).where(Route.codigo_ut == "542470"))
        stop = db.scalar(select(RouteStop).where(RouteStop.route_id == route.id))
        assert route.route_date == date(2026, 9, 14)
        assert stop.order_number == "14397867"
        assert stop.customer_address == "Rua Teste, 26, Centro"
